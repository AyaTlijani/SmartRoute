"""
orders_import.py — Read the daily Excel file and check every row.

Flow:
    read_excel(file)   → list of raw row dicts (headers normalised)
    validate_rows(...) → each row with cleaned data, errors, warnings
                         and "did you mean" suggestions

Rows with errors cannot be imported. Rows with warnings (possible
duplicate, risky customer, approximate location) can.
"""

import io
import re
import unicodedata
from typing import Any, Dict, List

from openpyxl import load_workbook

from tunisia import GOVERNORATES, approximate_location, match_governorate

# Template column → accepted header spellings (English, French, short forms).
COLUMNS = {
    "customer_name": ["customer_name", "customer", "name", "nom", "client", "nom client"],
    "phone": ["phone", "telephone", "tel", "mobile", "numero", "gsm"],
    "governorate": ["governorate", "gouvernorat", "region", "state"],
    "city": ["city", "ville", "delegation", "town"],
    "address": ["address", "adresse"],
    "landmark": ["landmark", "repere", "point de repere", "reference"],
    "product": ["product", "produit", "article", "item"],
    "cod_amount_dt": ["cod_amount_dt", "cod", "amount", "montant", "prix", "price", "cod amount"],
    "slot": ["slot", "creneau", "time slot", "time"],
    "notes": ["notes", "note", "remarque", "comment"],
    "lat": ["lat", "latitude"],
    "lng": ["lng", "lon", "long", "longitude"],
}

TEMPLATE_HEADERS = list(COLUMNS)
REQUIRED = ["customer_name", "phone", "governorate", "address", "cod_amount_dt"]


def _norm(text: Any) -> str:
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[\s_\-]+", " ", text.lower()).strip()


_HEADER_LOOKUP = {_norm(alias): col for col, aliases in COLUMNS.items() for alias in aliases}


def read_excel(content: bytes) -> List[Dict[str, Any]]:
    """Read the first sheet. Unknown columns are ignored."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        raise ValueError("The Excel file is empty.")
    mapping = {i: _HEADER_LOOKUP.get(_norm(h)) for i, h in enumerate(header)}
    if not {"phone", "governorate"} <= set(mapping.values()):
        raise ValueError(
            "Could not find the expected columns. Download the template "
            "and keep its header row."
        )
    out = []
    for values in it:
        if values is None or all(v in (None, "") for v in values):
            continue
        row = {col: "" for col in COLUMNS}
        for i, v in enumerate(values):
            col = mapping.get(i)
            if col:
                row[col] = "" if v is None else v
        out.append(row)
    return out


def normalize_phone(value: Any) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if isinstance(value, float) and str(value).endswith(".0"):
        digits = digits[:-1]
    for prefix in ("00216", "216"):
        if len(digits) > 8 and digits.startswith(prefix):
            digits = digits[len(prefix):]
    return digits


def parse_slot(value: Any):
    v = _norm(value)
    if v in ("", "any", "all day", "journee", "0"):
        return 0
    if v in ("1", "morning", "matin", "am", "1.0"):
        return 1
    if v in ("2", "afternoon", "apres midi", "pm", "2.0", "soir"):
        return 2
    return None


def parse_amount(value: Any):
    if isinstance(value, (int, float)):
        return float(value)
    v = re.sub(r"(dt|tnd|dinars?)", "", str(value or "").lower()).replace(",", ".").strip()
    try:
        return float(v)
    except ValueError:
        return None


def validate_rows(raw_rows: List[Dict[str, Any]], delivery_date: str, conn) -> List[Dict[str, Any]]:
    # Customer history by phone, for "risky customer" warnings.
    history = {
        r["phone"]: (r["total"], r["failed"])
        for r in conn.execute(
            "SELECT phone, COUNT(*) total, "
            "SUM(status = 'failed') failed FROM orders "
            "WHERE status IN ('delivered', 'failed') GROUP BY phone"
        )
    }
    already = {
        r["phone"]
        for r in conn.execute(
            "SELECT phone FROM orders WHERE delivery_date = ?", (delivery_date,)
        )
    }

    seen_phones: Dict[str, int] = {}
    results = []

    for i, raw in enumerate(raw_rows):
        row_no = i + 2  # Excel row number (header is row 1)
        errors, warnings, suggestions = [], [], {}
        clean_text = {k: str(raw.get(k, "") or "").strip() for k in COLUMNS}

        for field in REQUIRED:
            if not clean_text[field]:
                errors.append(f"Missing {field.replace('_', ' ')}")

        phone = normalize_phone(raw.get("phone"))
        if clean_text["phone"] and not re.fullmatch(r"[2-9]\d{7}", phone):
            errors.append(f"Invalid phone '{clean_text['phone']}' (need 8 digits)")

        governorate = None
        if clean_text["governorate"]:
            governorate, suggestion = match_governorate(clean_text["governorate"])
            if not governorate:
                if suggestion:
                    errors.append(f"Unknown governorate '{clean_text['governorate']}', did you mean {suggestion}?")
                    suggestions["governorate"] = suggestion
                else:
                    errors.append(f"Unknown governorate '{clean_text['governorate']}'")

        amount = parse_amount(raw.get("cod_amount_dt"))
        if clean_text["cod_amount_dt"] and (amount is None or amount < 0 or amount > 10000):
            errors.append(f"Invalid cash amount '{clean_text['cod_amount_dt']}'")

        slot = parse_slot(raw.get("slot"))
        if slot is None:
            errors.append(f"Unknown slot '{clean_text['slot']}' (use morning, afternoon or leave empty)")

        # Location: explicit lat/lng, else offline city/governorate table.
        lat = lng = None
        precision = None
        try:
            if clean_text["lat"] and clean_text["lng"]:
                lat, lng = float(raw["lat"]), float(raw["lng"])
                if not (30 <= lat <= 38 and 7 <= lng <= 12):
                    raise ValueError
                precision = "exact"
        except (TypeError, ValueError):
            errors.append("lat/lng are outside Tunisia")
            lat = lng = None
        if governorate and lat is None:
            lat, lng, precision = approximate_location(governorate, clean_text["city"])
            if precision == "governorate":
                warnings.append("Approximate location (governorate centre); driver should call")

        if phone and re.fullmatch(r"\d{8}", phone):
            if phone in seen_phones:
                warnings.append(f"Same phone as row {seen_phones[phone]}: possible duplicate")
            else:
                seen_phones[phone] = row_no
            if phone in already:
                warnings.append("Customer already has an order on this date")
            total, failed = history.get(phone, (0, 0))
            if failed and failed >= 2:
                warnings.append(f"Risky customer: {failed} of {total} past deliveries failed")

        results.append({
            "row": row_no,
            "raw": {k: clean_text[k] for k in COLUMNS},
            "data": {
                "customer_name": clean_text["customer_name"],
                "phone": phone,
                "governorate": governorate,
                "city": clean_text["city"],
                "address": clean_text["address"],
                "landmark": clean_text["landmark"],
                "product": clean_text["product"],
                "cod_amount": amount,
                "slot": slot,
                "notes": clean_text["notes"],
                "lat": lat,
                "lng": lng,
                "location_precision": precision,
            },
            "errors": errors,
            "warnings": warnings,
            "suggestions": suggestions,
            "status": "error" if errors else ("warning" if warnings else "ok"),
        })

    return results


def insert_orders(conn, delivery_date: str, rows: List[Dict[str, Any]]) -> int:
    cols = ["customer_name", "phone", "governorate", "city", "address", "landmark",
            "product", "cod_amount", "slot", "notes", "lat", "lng", "location_precision"]
    conn.executemany(
        f"INSERT INTO orders (delivery_date, {', '.join(cols)}) "
        f"VALUES (?, {', '.join('?' for _ in cols)})",
        [[delivery_date] + [r.get(c) for c in cols] for r in rows],
    )
    return len(rows)


assert set(REQUIRED) <= set(COLUMNS) and len(GOVERNORATES) == 24
