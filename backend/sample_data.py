"""
sample_data.py — Realistic Tunisian demo data.

    seed_demo(conn)          drivers + 3 weeks of past deliveries
    sample_orders_xlsx(day)  a day's confirmed-orders Excel file,
                             with a few deliberate mistakes to catch
    template_xlsx()          the empty template admins fill in
"""

import io
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from auth import hash_password
from orders_import import TEMPLATE_HEADERS
from tunisia import CITIES

HUB = ("SmartRoute Hub, Charguia", 36.8400, 10.2050)

DRIVERS = [
    # name, phone, vehicle, zones, depot (name, lat, lng)
    ("Ahmed Trabelsi", "98123456", "scooter", "Tunis,Ariana", HUB),
    ("Sami Ben Ali", "55234567", "scooter", "Tunis,Ben Arous", HUB),
    ("Youssef Gharbi", "22345678", "car", "Ariana,Manouba,Ben Arous", HUB),
    ("Mehdi Jlassi", "97456789", "car", "Sousse,Monastir,Mahdia", ("Sousse relay point", 35.8256, 10.6084)),
    ("Karim Masmoudi", "50567890", "car", "Sfax", ("Sfax relay point", 34.7406, 10.7603)),
]

FIRST = ["Amira", "Mohamed", "Fatma", "Ahmed", "Yasmine", "Oussema", "Sarra", "Khalil",
         "Nour", "Aymen", "Rania", "Hamza", "Ines", "Skander", "Mariem", "Wassim",
         "Salma", "Bilel", "Eya", "Firas", "Olfa", "Hichem", "Chaima", "Seif"]
LAST = ["Ben Salah", "Trabelsi", "Jaziri", "Gharbi", "Hammami", "Mansouri", "Ayari",
        "Chaabane", "Dridi", "Ferchichi", "Khemiri", "Mejri", "Saidi", "Zouari",
        "Baccouche", "Chebbi", "Jebali", "Kefi", "Brahmi", "Masmoudi"]
STREETS = ["Rue de Marseille", "Avenue Habib Bourguiba", "Rue Ibn Khaldoun",
           "Avenue de la Liberte", "Rue de Palestine", "Avenue Hedi Nouira",
           "Rue 18 Janvier", "Rue Farhat Hached", "Avenue Taieb Mhiri", "Rue Mongi Slim"]
LANDMARKS = ["Near the pharmacy", "Behind the mosque", "Next to Monoprix",
             "In front of the school", "Above the bakery", "Blue door, 2nd floor",
             "Near the Total station", "Opposite the post office", ""]
PRODUCTS = ["Sneakers size 42", "Black hoodie L", "Phone case", "Perfume 100ml",
            "Leather handbag", "Kids dress", "Smart watch", "Hair dryer", "Makeup kit",
            "Bluetooth earbuds", "Abaya", "Kitchen set"]
FAIL_REASONS = (["No answer"] * 8 + ["Customer refused"] * 5 + ["Not at home"] * 4
                + ["Address not found"] * 2 + ["Other"])

# Demo logins (shown on the login page; this is a local demo, not production).
ADMIN_LOGIN = ("admin", "admin123", "Leila Haddad")
DRIVER_PASSWORD = "driver123"

# Where orders come from: Grand Tunis dominates, a few from far cities
# that no driver covers (they show up as "unassigned").
CITY_WEIGHTS = (
    [(c, 6) for c, (g, *_ ) in CITIES.items() if g in ("Tunis", "Ariana", "Ben Arous", "Manouba")]
    + [(c, 8) for c, (g, *_ ) in CITIES.items() if g in ("Sousse", "Monastir")]
    + [("Sfax Centre", 14), ("Sakiet Ezzit", 9), ("Hammamet", 1), ("Nabeul Centre", 1),
       ("Djerba Houmt Souk", 1), ("Bizerte Centre", 1)]
)
GOV_FAIL_BONUS = {"Manouba": 0.10, "Ben Arous": 0.05}


def _phone(rng):
    return rng.choice("2459") + "".join(rng.choice("0123456789") for _ in range(7))


def _risky_phones():
    rng = random.Random(7)
    return [_phone(rng) for _ in range(12)]


def _name_for(phone):
    """Same phone number → same customer name across days."""
    r = random.Random(phone)
    return f"{r.choice(FIRST)} {r.choice(LAST)}"


def _make_order(rng, phone=None):
    city = rng.choices([c for c, _ in CITY_WEIGHTS], [w for _, w in CITY_WEIGHTS])[0]
    gov, lat, lng = CITIES[city]
    phone = phone or _phone(rng)
    return {
        "customer_name": _name_for(phone),
        "phone": phone,
        "governorate": gov,
        "city": city,
        "address": f"{rng.randint(1, 120)} {rng.choice(STREETS)}",
        "landmark": rng.choice(LANDMARKS),
        "product": rng.choice(PRODUCTS),
        "cod_amount_dt": round(rng.uniform(25, 180) * 2) / 2,
        "slot": rng.choices(["morning", "afternoon", ""], [1, 1, 2])[0],
        "notes": rng.choice(["", "", "", "Call before arriving", "Leave with the guard"]),
        # ~1.5 km spread around the city centre, like a shared GPS pin.
        "lat": round(lat + rng.uniform(-0.012, 0.012), 5),
        "lng": round(lng + rng.uniform(-0.012, 0.012), 5),
    }


def seed_demo(conn, days: int = 21, live_today: bool = None) -> dict:
    """
    Wipe the database and load the demo company (see demo_seed.py).
    live_today=False (or SMARTROUTE_SEED_TODAY=0) skips today's in-progress
    day and tomorrow's orders; the tests use that to start from a clean day.
    """
    import os
    import demo_seed
    if live_today is None:
        live_today = os.getenv("SMARTROUTE_SEED_TODAY", "1") != "0"
    return demo_seed.seed(conn, DRIVERS, ADMIN_LOGIN, DRIVER_PASSWORD, _make_order, _risky_phones, _phone,
                          days=days, live_today=live_today)


def ensure_accounts(conn) -> None:
    """
    Databases created before accounts existed have drivers but no users.
    Give them the demo admin and one login per driver (first name).
    """
    if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]:
        return
    conn.execute("INSERT INTO users (username, password_hash, role, full_name) VALUES (?, ?, 'admin', ?)",
                 (ADMIN_LOGIN[0], hash_password(ADMIN_LOGIN[1]), ADMIN_LOGIN[2]))
    for d in conn.execute("SELECT id, name FROM drivers").fetchall():
        username = d["name"].split()[0].lower()
        if conn.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
            username = f"{username}{d['id']}"
        conn.execute("INSERT INTO users (username, password_hash, role, driver_id, full_name) "
                     "VALUES (?, ?, 'driver', ?, ?)",
                     (username, hash_password(DRIVER_PASSWORD), d["id"], d["name"]))


def _workbook(rows, title="Orders"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E40AF")
    for r in rows:
        ws.append([r.get(h, "") for h in TEMPLATE_HEADERS])
    widths = {"customer_name": 22, "address": 28, "landmark": 24, "product": 20, "notes": 22}
    for i, h in enumerate(TEMPLATE_HEADERS):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = widths.get(h, 14)
    ws.freeze_panes = "A2"
    return wb


def _to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_xlsx() -> bytes:
    wb = _workbook([
        {"customer_name": "Amira Ben Salah", "phone": "22345678", "governorate": "Ariana",
         "city": "La Soukra", "address": "12 Rue des Roses", "landmark": "Near the pharmacy",
         "product": "Sneakers size 41", "cod_amount_dt": 89.5, "slot": "morning",
         "notes": "Call before arriving"},
    ])
    info = wb.create_sheet("How to fill")
    for line in [
        ["Required: customer_name, phone, governorate, address, cod_amount_dt"],
        ["phone: 8 digits, +216 prefix is fine"],
        ["governorate: one of the 24 (Tunis, Ariana, Ben Arous, Sfax, ...)"],
        ["slot: morning (9-12), afternoon (13-17) or empty for any time"],
        ["lat / lng: optional GPS pin; without it the city or governorate centre is used"],
        ["Delete the example row before uploading."],
    ]:
        info.append(line)
    info.column_dimensions["A"].width = 80
    return _to_bytes(wb)


def sample_orders_xlsx(day: date, n: int = 58) -> bytes:
    """A realistic day of confirmed orders, including 5 rows to catch."""
    rng = random.Random(day.toordinal())
    risky = _risky_phones()
    rows = [_make_order(rng, rng.choice(risky) if rng.random() < 0.07 else None) for _ in range(n)]
    for r in rows[::5]:
        r["lat"] = r["lng"] = ""  # some rows without a GPS pin → city lookup

    # Deliberate problems the validator should flag:
    rows[3]["phone"] = "2234 567"             # 7 digits
    rows[8]["governorate"] = "Ariena"          # typo → "did you mean Ariana?"
    rows[8]["lat"] = rows[8]["lng"] = ""
    rows[8]["city"] = "Ennasr"
    rows[13]["address"] = ""                   # missing address
    rows[17]["cod_amount_dt"] = "cent dinars"  # not a number
    dup = dict(rows[5]); dup["product"] = rng.choice(PRODUCTS)
    rows.insert(20, dup)                       # same customer twice
    rows[25]["phone"] = "+216 " + rows[25]["phone"]  # valid, prefix stripped
    rows[30]["governorate"] = "Benarous"       # alias → Ben Arous
    for i, place in ((40, "Cite El Khadra"), (45, "Borj Louzir")):
        rows[i]["city"] = place                # not in the city table →
        rows[i]["lat"] = rows[i]["lng"] = ""   # governorate-level location
    return _to_bytes(_workbook(rows))
